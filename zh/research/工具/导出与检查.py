"""Export paper manuscripts to Word and check document package integrity.

This utility verifies document structure, not scientific validity or page layout.
Run with the bundled Python runtime and an installed Pandoc executable.
"""
from __future__ import annotations

import argparse
import hashlib
import json
import posixpath
import re
import shutil
import subprocess
import tempfile
from pathlib import Path
from zipfile import ZipFile

from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor
from lxml import etree


ROOT = Path(__file__).resolve().parents[1]
W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def body_text(source: str) -> str:
    return re.sub(r"\A---\r?\n.*?\r?\n---\r?\n", "", source, count=1, flags=re.S).lstrip()


def package_errors(path: Path) -> list[str]:
    errors = []
    with ZipFile(path) as archive:
        names = set(archive.namelist())
        if archive.testzip() is not None:
            errors.append("ZIP CRC failure")
        for relfile in sorted(n for n in names if n.endswith(".rels")):
            source = "" if relfile == "_rels/.rels" else posixpath.join(
                posixpath.dirname(posixpath.dirname(relfile)), posixpath.basename(relfile)[:-5]
            )
            rels = {}
            for element in etree.fromstring(archive.read(relfile)):
                rid = element.get("Id")
                rels[rid] = element
                if element.get("TargetMode") == "External":
                    continue
                target = posixpath.normpath(posixpath.join(
                    posixpath.dirname(source), element.get("Target", "")
                )).lstrip("/")
                if target not in names:
                    errors.append(f"{source}: missing relationship target {target}")
            if source in names and source.endswith(".xml"):
                for element in etree.fromstring(archive.read(source)).iter():
                    for attr in ("id", "embed", "link"):
                        rid = element.get(f"{{{R}}}{attr}")
                        if rid and rid not in rels:
                            errors.append(f"{source}: unresolved relationship {rid}")
        defined = {
            e.get(f"{{{W}}}styleId")
            for e in etree.fromstring(archive.read("word/styles.xml")).iter(f"{{{W}}}style")
        }
        doc_root = etree.fromstring(archive.read("word/document.xml"))
        for tag in ("pStyle", "rStyle", "tblStyle"):
            for element in doc_root.iter(f"{{{W}}}{tag}"):
                style = element.get(f"{{{W}}}val")
                if style not in defined:
                    errors.append(f"undefined style {style}")
    return errors


def reference_document(pandoc: str, target: Path) -> None:
    target.write_bytes(subprocess.check_output([pandoc, "--print-default-data-file", "reference.docx"]))
    document = Document(target)
    section = document.sections[0]
    section.page_width, section.page_height = Cm(21), Cm(29.7)
    section.top_margin, section.bottom_margin = Cm(2.3), Cm(2.3)
    section.left_margin, section.right_margin = Cm(2.5), Cm(2.5)
    for name in ("Normal", "Body Text", "First Paragraph", "Compact", "Block Text"):
        if name not in document.styles:
            continue
        style = document.styles[name]
        style.font.name, style.font.size = "Times New Roman", Pt(11)
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Songti SC")
        style.paragraph_format.line_spacing = 1.3
        style.paragraph_format.space_after = Pt(6)
    for name, size in (("Title", 19), ("Heading 1", 16), ("Heading 2", 13), ("Heading 3", 11.5)):
        style = document.styles[name]
        style.font.name, style.font.size = "Arial", Pt(size)
        style.font.color.rgb = RGBColor.from_string("172B4D")
        style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Heiti SC")
        style.paragraph_format.keep_with_next = True
    footer = section.footer.paragraphs[0]
    footer.alignment = 1
    footer.add_run("研究初稿  |  ")
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    footer._p.append(field)
    document.save(target)


def audit(source: Path, word: Path) -> dict:
    raw = source.read_text(encoding="utf-8")
    body = body_text(raw)
    document = Document(word)
    paragraphs = "\n".join(p.text for p in document.paragraphs)
    errors = package_errors(word)
    refs = re.findall(r"^\[(\d+)\]\s", body, re.M)
    if refs and [int(x) for x in refs] != list(range(1, len(refs) + 1)):
        errors.append("reference numbers are not sequential")
    for number in refs:
        if f"[{number}]" not in paragraphs:
            errors.append(f"missing Word reference [{number}]")
    abs_match = re.search(r"^##\s+(?:\d+[.、 ]+)?Abstract\s*\n(.*?)(?=^## |\*\*Keywords|\*\*Key words|\*\*关键词|^Keywords:)", body, re.M | re.S | re.I)
    abstract_words = None
    if abs_match:
        abstract_words = len(re.findall(r"[A-Za-z]+(?:[-'][A-Za-z]+)*", abs_match.group(1)))
    warnings = []
    if abstract_words is None or not 150 <= abstract_words <= 300:
        warnings.append(f"English abstract needs inspection: {abstract_words}")
    for label, target in re.findall(r"\[([^\]]+)\]\(([^)]+)\)", body):
        if target.startswith(("http:", "https:", "#", "mailto:")):
            continue
        local = target.strip("<>").split("#", 1)[0]
        if local and not (source.parent / local).exists():
            warnings.append(f"missing local link: {label} -> {target}")
    return {
        "paper": str(source.relative_to(ROOT)),
        "markdown_sha256": hashlib.sha256(source.read_bytes()).hexdigest(),
        "word_sha256": hashlib.sha256(word.read_bytes()).hexdigest(),
        "chinese_characters": len(re.findall(r"[\u4e00-\u9fff]", body)),
        "english_abstract_words": abstract_words,
        "reference_count": len(refs),
        "paragraphs": len(document.paragraphs),
        "tables": [{"rows": len(t.rows), "columns": len(t.columns)} for t in document.tables],
        "math_objects": document._element.xml.count("<m:oMath>"),
        "errors": errors,
        "warnings": warnings,
    }


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--check-only", action="store_true")
    parser.add_argument("--paper", help="Optional paper directory prefix, for example 02")
    args = parser.parse_args()
    papers = sorted(ROOT.glob("[0-9][0-9]-*/论文初稿.md"))
    if args.paper:
        papers = [p for p in papers if p.parent.name.startswith(args.paper)]
    pandoc = shutil.which("pandoc")
    if not args.check_only and not pandoc:
        raise SystemExit("Pandoc is required for export; no installation was attempted.")
    reports = []
    with tempfile.TemporaryDirectory(prefix="paper_series_") as directory:
        temp = Path(directory)
        reference = temp / "reference.docx"
        if not args.check_only:
            reference_document(pandoc, reference)
        for source in papers:
            word = source.with_suffix(".docx")
            if not args.check_only:
                intermediate = temp / "paper.md"
                intermediate.write_text(body_text(source.read_text(encoding="utf-8")), encoding="utf-8")
                built = temp / "paper.docx"
                result = subprocess.run([
                    pandoc, str(intermediate), "-f", "markdown+tex_math_single_backslash", "-t", "docx",
                    "--reference-doc", str(reference), "-o", str(built)
                ], capture_output=True, text=True, check=True)
                if result.stderr.strip():
                    raise RuntimeError(f"Review Pandoc warnings before export: {result.stderr}")
                errors = package_errors(built)
                if errors:
                    raise RuntimeError(errors)
                shutil.copyfile(built, word)
            reports.append(audit(source, word))
    output = ROOT / ("文档结构检查.json" if not args.paper else f"文档结构检查-{args.paper}.json")
    output.write_text(json.dumps(reports, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(reports, ensure_ascii=False, indent=2))
    if any(report["errors"] for report in reports):
        raise SystemExit(1)


if __name__ == "__main__":
    main()
